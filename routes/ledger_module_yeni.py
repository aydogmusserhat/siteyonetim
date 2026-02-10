# routes/ledger_module.py
# Admin + SuperAdmin: Daire bazında aylık borç/ödeme ekstre ekranı + Excel/PDF export
# ✅ "Bu Daire – Excel/PDF" tıklanınca İşlemler içindeki TÜM detaylar da export'a yazılır:
#    - Dönem Borçları
#    - Dönem Ödemeleri
#    - Devreden Borçlar (dönem başlamadan önceki açık/partial borçların kalan tutarı)
#    - Devreden Ödemeler (bu dönem içinde önceki dönem borçlarına yapılan ödemeler)

import io
from datetime import date, timedelta, datetime
from decimal import Decimal
from collections import defaultdict
from typing import Optional

from flask import request, render_template, redirect, url_for, flash, session, send_file
from sqlalchemy import and_, or_, func

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from models import db
from models.user_model import User
from models.apartment_model import Apartment
from models.bill_model import Bill
from models.payment_model import Payment
from models.site_model import Site


def register_ledger_routes(admin_bp, admin_required, _get_current_admin, _parse_date_flex):
    """
    admin_routes.py içinde 1 kez çağır.
    Route'lar admin_bp'ye eklenir.
    """

    # -------------------------
    # Helpers
    # -------------------------
    def _month_range_from_str(month_str: str):
        month_str = (month_str or "").strip()
        if not month_str:
            return None, None
        y, m = month_str.split("-")
        y = int(y)
        m = int(m)
        start = date(y, m, 1)
        if m == 12:
            end = date(y + 1, 1, 1)
        else:
            end = date(y, m + 1, 1)
        return start, end

    def _resolve_period_from_request():
        """
        Öncelik:
          1) month=YYYY-MM
          2) from=... to=... (esnek parse)
          3) yoksa bu ay
        """
        month = (request.args.get("month") or "").strip()
        if month:
            s, e = _month_range_from_str(month)
            return month, s, e

        date_from = (request.args.get("from") or "").strip()
        date_to = (request.args.get("to") or "").strip()

        if date_from and date_to:
            s = _parse_date_flex(date_from)
            e_inclusive = _parse_date_flex(date_to)
            e = e_inclusive + timedelta(days=1)  # inclusive -> exclusive
            return "", s, e

        # default this month
        today = date.today()
        s = date(today.year, today.month, 1)
        if today.month == 12:
            e = date(today.year + 1, 1, 1)
            month_label = f"{today.year}-12"
        else:
            e = date(today.year, today.month + 1, 1)
            month_label = f"{today.year}-{today.month:02d}"
        return month_label, s, e

    def _get_site_scope_for_admin(admin_user: User):
        """
        Admin => site zorunlu
        Super admin => aktif site seçiliyse o site, değilse global
        """
        site_id = session.get("active_site_id") or (admin_user.site_id if getattr(admin_user, "site_id", None) else None)
        is_super = (getattr(admin_user, "role", "") == "super_admin")
        global_mode = is_super and not site_id
        return site_id, global_mode

    def _ledger_build_dataset(
        site_id: Optional[int],
        global_mode: bool,
        start: date,
        end: date,
        apartment_id: Optional[int],
        q: Optional[str]
    ):
        """
        start-end aralığında:
          - billed: dönem borç toplamı
          - paid: dönem ödeme toplamı
          - open_balance: genel açık bakiye (open/partial)
          - bills/payments: dönem detay listeleri
          ✅ - carry_bills: dönem başlamadan önceki açık/partial borçlar (kalanlarıyla)
          ✅ - carry_payments: bu dönem içinde önceki dönem borçlarına yapılan ödemeler
          ✅ - carry_paid: carry_payments toplamı (ekrandaki "Devreden Ödeme" sütunu için)
        """

        # --- daire temel sorgusu ---
        apt_q = Apartment.query
        if not global_mode:
            apt_q = apt_q.filter(Apartment.site_id == site_id)

        if apartment_id:
            apt_q = apt_q.filter(Apartment.id == apartment_id)

        if q:
            like = f"%{q}%"
            # Bu alanlar sende varsa çalışır: block/floor/number/owner_name
            apt_q = apt_q.filter(
                (Apartment.block.ilike(like)) |
                (Apartment.floor.ilike(like)) |
                (Apartment.number.ilike(like)) |
                (Apartment.owner_name.ilike(like))
            )

        apartments = apt_q.order_by(
            Apartment.block.asc(),
            Apartment.floor.asc(),
            Apartment.number.asc()
        ).all()

        # --- her bill için toplam ödenen subquery (tüm zamanlar) ---
        pay_sum_subq = (
            db.session.query(
                Payment.bill_id.label("bill_id"),
                func.coalesce(func.sum(Payment.amount), 0).label("paid_sum"),
            )
            .group_by(Payment.bill_id)
            .subquery()
        )

        # --- billed (period) ---
        billed_q = db.session.query(
            Bill.apartment_id.label("apartment_id"),
            func.coalesce(func.sum(Bill.amount), 0).label("billed_sum")
        )

        if not global_mode:
            billed_q = billed_q.filter(Bill.site_id == site_id)

        billed_q = billed_q.filter(
            or_(
                and_(Bill.due_date.isnot(None), Bill.due_date >= start, Bill.due_date < end),
                and_(Bill.due_date.is_(None), Bill.created_at >= start, Bill.created_at < end),
            )
        ).group_by(Bill.apartment_id)

        billed_map = {r.apartment_id: r.billed_sum for r in billed_q.all()}

        # --- paid (period) ---
        paid_q = db.session.query(
            Payment.apartment_id.label("apartment_id"),
            func.coalesce(func.sum(Payment.amount), 0).label("paid_sum")
        )

        if not global_mode:
            paid_q = paid_q.filter(Payment.site_id == site_id)

        paid_q = paid_q.filter(
            Payment.payment_date >= start,
            Payment.payment_date < end,
        ).group_by(Payment.apartment_id)

        paid_map = {r.apartment_id: r.paid_sum for r in paid_q.all()}

        # --- carry totals (start öncesi toplam borç/ödeme) ---
        start_dt = datetime.combine(start, datetime.min.time())

        carry_billed_q = db.session.query(
            Bill.apartment_id.label("apartment_id"),
            func.coalesce(func.sum(Bill.amount), 0).label("billed_sum"),
        )
        if not global_mode:
            carry_billed_q = carry_billed_q.filter(Bill.site_id == site_id)
        carry_billed_q = carry_billed_q.filter(
            or_(
                and_(Bill.due_date.isnot(None), Bill.due_date < start),
                and_(Bill.due_date.is_(None), Bill.created_at < start_dt),
            )
        ).group_by(Bill.apartment_id)
        carry_billed_map = {r.apartment_id: r.billed_sum for r in carry_billed_q.all()}

        carry_paid_q = db.session.query(
            Payment.apartment_id.label("apartment_id"),
            func.coalesce(func.sum(Payment.amount), 0).label("paid_sum"),
        )
        if not global_mode:
            carry_paid_q = carry_paid_q.filter(Payment.site_id == site_id)
        carry_paid_q = carry_paid_q.filter(Payment.payment_date < start).group_by(Payment.apartment_id)
        carry_paid_map = {r.apartment_id: r.paid_sum for r in carry_paid_q.all()}

        # --- open balance (overall) ---
        open_q = (
            db.session.query(
                Bill.apartment_id.label("apartment_id"),
                func.coalesce(
                    func.sum(Bill.amount - func.coalesce(pay_sum_subq.c.paid_sum, 0)),
                    0
                ).label("open_sum")
            )
            .outerjoin(pay_sum_subq, pay_sum_subq.c.bill_id == Bill.id)
            .filter(Bill.status.in_(["open", "partial"]))
        )

        if not global_mode:
            open_q = open_q.filter(Bill.site_id == site_id)

        open_q = open_q.group_by(Bill.apartment_id)
        open_map = {r.apartment_id: r.open_sum for r in open_q.all()}

        # --- detaylar (period içinde borç/ödeme listesi) ---
        bill_detail_q = db.session.query(Bill).join(Apartment, Bill.apartment_id == Apartment.id)
        if not global_mode:
            bill_detail_q = bill_detail_q.filter(Bill.site_id == site_id)
        bill_detail_q = bill_detail_q.filter(
            or_(
                and_(Bill.due_date.isnot(None), Bill.due_date >= start, Bill.due_date < end),
                and_(Bill.due_date.is_(None), Bill.created_at >= start, Bill.created_at < end),
            )
        )
        if apartment_id:
            bill_detail_q = bill_detail_q.filter(Bill.apartment_id == apartment_id)

        bills = bill_detail_q.order_by(Bill.due_date.desc().nullslast(), Bill.created_at.desc()).all()
        bills_by_apt = defaultdict(list)
        for b in bills:
            bills_by_apt[b.apartment_id].append(b)

        pay_detail_q = db.session.query(Payment).join(Apartment, Payment.apartment_id == Apartment.id)
        if not global_mode:
            pay_detail_q = pay_detail_q.filter(Payment.site_id == site_id)
        pay_detail_q = pay_detail_q.filter(Payment.payment_date >= start, Payment.payment_date < end)
        if apartment_id:
            pay_detail_q = pay_detail_q.filter(Payment.apartment_id == apartment_id)

        payments = pay_detail_q.order_by(Payment.payment_date.desc()).all()
        pays_by_apt = defaultdict(list)
        for p in payments:
            pays_by_apt[p.apartment_id].append(p)

        # =========================================================
        # ✅ DEVREDEN DETAYLAR
        # =========================================================

        # 1) start öncesine kadar yapılmış ödemeler (bill bazında) => kalan hesaplamak için
        pay_before_start_subq = (
            db.session.query(
                Payment.bill_id.label("bill_id"),
                func.coalesce(func.sum(Payment.amount), 0).label("paid_sum"),
            )
            .filter(Payment.payment_date < start)
            .group_by(Payment.bill_id)
            .subquery()
        )

        # 2) Devreden borçlar: start öncesi bill'ler ve open/partial olanlar, kalan>0 ise listele
        carry_bill_q = (
            db.session.query(
                Bill,
                (Bill.amount - func.coalesce(pay_before_start_subq.c.paid_sum, 0)).label("remaining")
            )
            .outerjoin(pay_before_start_subq, pay_before_start_subq.c.bill_id == Bill.id)
            .filter(Bill.status.in_(["open", "partial"]))
            .filter(
                or_(
                    and_(Bill.due_date.isnot(None), Bill.due_date < start),
                    and_(Bill.due_date.is_(None), Bill.created_at < start_dt),
                )
            )
        )
        if not global_mode:
            carry_bill_q = carry_bill_q.filter(Bill.site_id == site_id)
        if apartment_id:
            carry_bill_q = carry_bill_q.filter(Bill.apartment_id == apartment_id)

        carry_bill_q = carry_bill_q.order_by(Bill.due_date.desc().nullslast(), Bill.created_at.desc())
        carry_bills_by_apt = defaultdict(list)
        for b, remaining in carry_bill_q.all():
            rem = Decimal(str(remaining or 0))
            if rem > 0:
                # template/export için b._remaining kullanacağız
                try:
                    setattr(b, "_remaining", float(rem))
                except Exception:
                    pass
                carry_bills_by_apt[b.apartment_id].append(b)

        # 3) Devreden ödemeler: bu dönem içindeki ödemeler ama bill'i start öncesi olanlar
        carry_pay_q = (
            db.session.query(Payment)
            .join(Bill, Payment.bill_id == Bill.id)
            .filter(Payment.payment_date >= start, Payment.payment_date < end)
            .filter(
                or_(
                    and_(Bill.due_date.isnot(None), Bill.due_date < start),
                    and_(Bill.due_date.is_(None), Bill.created_at < start_dt),
                )
            )
        )
        if not global_mode:
            carry_pay_q = carry_pay_q.filter(Payment.site_id == site_id)
        if apartment_id:
            carry_pay_q = carry_pay_q.filter(Payment.apartment_id == apartment_id)

        carry_pay_q = carry_pay_q.order_by(Payment.payment_date.desc())
        carry_pays_by_apt = defaultdict(list)
        carry_paid_sum_by_apt = defaultdict(Decimal)

        for p in carry_pay_q.all():
            carry_pays_by_apt[p.apartment_id].append(p)
            carry_paid_sum_by_apt[p.apartment_id] += Decimal(str(getattr(p, "amount", 0) or 0))

        # --- output rows ---
        rows = []
        for apt in apartments:
            billed = Decimal(str(billed_map.get(apt.id, 0) or 0))
            paid = Decimal(str(paid_map.get(apt.id, 0) or 0))
            open_bal = Decimal(str(open_map.get(apt.id, 0) or 0))

            carry_billed_total = Decimal(str(carry_billed_map.get(apt.id, 0) or 0))
            carry_paid_total_before = Decimal(str(carry_paid_map.get(apt.id, 0) or 0))
            carry_net = carry_billed_total - carry_paid_total_before

            carry_debt = max(carry_net, Decimal("0"))
            carry_credit = max(-carry_net, Decimal("0"))

            carry_paid_in_period = Decimal(str(carry_paid_sum_by_apt.get(apt.id, 0) or 0))

            rows.append({
                "apartment": apt,
                "site_id": getattr(apt, "site_id", None),
                "billed": billed,
                "paid": paid,
                "net": paid - billed,
                "open_balance": open_bal,

                # ekrandaki özet için
                "carry_debt": carry_debt,
                "carry_credit": carry_credit,
                "carry_paid": carry_paid_in_period,  # ✅ ekrandaki "Devreden Ödeme"

                # ekran "işlemler" detay tabloları için
                "bills": bills_by_apt.get(apt.id, []),
                "payments": pays_by_apt.get(apt.id, []),
                "carry_bills": carry_bills_by_apt.get(apt.id, []),
                "carry_payments": carry_pays_by_apt.get(apt.id, []),
            })

        site_map = {}
        if global_mode:
            sites = Site.query.order_by(Site.name.asc()).all()
            site_map = {s.id: s.name for s in sites}

        return rows, site_map

    # -------------------------
    # Routes
    # -------------------------
    @admin_bp.route("/ledger", methods=["GET"])
    @admin_required
    def apartment_ledger():
        admin_user = _get_current_admin()
        if not admin_user:
            flash("Kullanıcı bulunamadı. Lütfen tekrar giriş yapın.", "error")
            return redirect(url_for("auth.logout"))

        site_id, global_mode = _get_site_scope_for_admin(admin_user)
        if not global_mode and not site_id:
            flash("Bu ekran için bir siteye atanmış olmanız gerekiyor.", "error")
            return redirect(url_for("admin.dashboard"))

        month_label, start, end = _resolve_period_from_request()
        apartment_id = request.args.get("apartment_id", type=int)
        q = (request.args.get("q") or "").strip() or None

        rows, site_map = _ledger_build_dataset(site_id, global_mode, start, end, apartment_id, q)

        apt_list_q = Apartment.query
        if not global_mode:
            apt_list_q = apt_list_q.filter_by(site_id=site_id)
        apt_list = apt_list_q.order_by(Apartment.block.asc(), Apartment.floor.asc(), Apartment.number.asc()).all()

        args = request.args.to_dict(flat=True)
        period_end_inclusive = end - timedelta(days=1)

        return render_template(
            "admin/ledger.html",
            args=args,
            period_end_inclusive=period_end_inclusive,
            admin_user=admin_user,
            global_mode=global_mode,
            month_label=month_label,
            period_start=start,
            period_end=end,
            rows=rows,
            site_map=site_map,
            apartments=apt_list,
            selected_apartment_id=apartment_id,
            q=q or "",
        )

    @admin_bp.route("/ledger/export.xlsx", methods=["GET"])
    @admin_required
    def apartment_ledger_export_xlsx():
        """
        ✅ Eğer apartment_id gelirse (Bu Daire – Excel), sadece o daire export olur
        ✅ İşlemler içindeki tüm detaylar ayrı sheet'lerde yazılır
        """
        admin_user = _get_current_admin()
        site_id, global_mode = _get_site_scope_for_admin(admin_user)

        if not global_mode and not site_id:
            flash("Bu işlem için bir siteye atanmış olmanız gerekiyor.", "error")
            return redirect(url_for("admin.dashboard"))

        month_label, start, end = _resolve_period_from_request()
        apartment_id = request.args.get("apartment_id", type=int)
        q = (request.args.get("q") or "").strip() or None

        rows, site_map = _ledger_build_dataset(site_id, global_mode, start, end, apartment_id, q)

        def _fmt_dt(d):
            if not d:
                return ""
            try:
                dd = getattr(d, "date", lambda: d)()
                return dd.strftime("%Y-%m-%d")
            except Exception:
                return str(d)

        wb = Workbook()

        # -------------------------
        # 1) Özet sheet
        # -------------------------
        ws = wb.active
        ws.title = "Özet"

        headers = []
        if global_mode:
            headers.append("Site")
        headers += [
            "Daire",
            "Malik",
            "Devreden Borç",
            "Devreden Ödeme",
            "Dönem Borç",
            "Dönem Ödeme",
            "Net (Ödeme-Borç)",
            "Açık Bakiye (Genel)",
        ]
        ws.append(headers)

        for r in rows:
            apt = r["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            owner = getattr(apt, "owner_name", "") or ""

            line = []
            if global_mode:
                line.append(site_map.get(r["site_id"], ""))
            line += [
                apt_label,
                owner,
                float(r.get("carry_debt", 0) or 0),
                float(r.get("carry_paid", 0) or 0),
                float(r["billed"]),
                float(r["paid"]),
                float(r["net"]),
                float(r["open_balance"]),
            ]
            ws.append(line)

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22

        # -------------------------
        # 2) Detay - Dönem Borçları
        # -------------------------
        ws_b = wb.create_sheet("Detay-Borçlar")
        b_headers = []
        if global_mode:
            b_headers.append("Site")
        b_headers += ["Daire", "Malik", "Tarih", "Açıklama", "Tutar", "Durum", "Bill ID"]
        ws_b.append(b_headers)

        for r in rows:
            apt = r["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            owner = getattr(apt, "owner_name", "") or ""
            for b in (r.get("bills") or []):
                b_date = getattr(b, "due_date", None) or getattr(b, "created_at", None)
                desc = getattr(b, "description", None) or getattr(b, "desc", None) or ""
                status = getattr(b, "status", "") or ""
                amount = float(getattr(b, "amount", 0) or 0)

                line = []
                if global_mode:
                    line.append(site_map.get(r["site_id"], ""))
                line += [apt_label, owner, _fmt_dt(b_date), desc, amount, status, getattr(b, "id", None)]
                ws_b.append(line)

        for col in range(1, len(b_headers) + 1):
            ws_b.column_dimensions[get_column_letter(col)].width = 22
        ws_b.column_dimensions[get_column_letter(4 if global_mode else 3)].width = 14
        ws_b.column_dimensions[get_column_letter(5 if global_mode else 4)].width = 48

        # -------------------------
        # 3) Detay - Dönem Ödemeleri
        # -------------------------
        ws_p = wb.create_sheet("Detay-Ödemeler")
        p_headers = []
        if global_mode:
            p_headers.append("Site")
        p_headers += ["Daire", "Malik", "Tarih", "Yöntem", "Tutar", "Bill ID", "Payment ID"]
        ws_p.append(p_headers)

        for r in rows:
            apt = r["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            owner = getattr(apt, "owner_name", "") or ""
            for p in (r.get("payments") or []):
                p_date = getattr(p, "payment_date", None)
                method = getattr(p, "method", None) or ""
                amount = float(getattr(p, "amount", 0) or 0)

                line = []
                if global_mode:
                    line.append(site_map.get(r["site_id"], ""))
                line += [apt_label, owner, _fmt_dt(p_date), method, amount, getattr(p, "bill_id", None), getattr(p, "id", None)]
                ws_p.append(line)

        for col in range(1, len(p_headers) + 1):
            ws_p.column_dimensions[get_column_letter(col)].width = 20
        ws_p.column_dimensions[get_column_letter(4 if global_mode else 3)].width = 14

        # -------------------------
        # 4) ✅ Devreden - Borçlar
        # -------------------------
        ws_cb = wb.create_sheet("Devreden-Borçlar")
        cb_headers = []
        if global_mode:
            cb_headers.append("Site")
        cb_headers += ["Daire", "Malik", "Tarih", "Açıklama", "Kalan", "Durum", "Bill ID"]
        ws_cb.append(cb_headers)

        for r in rows:
            apt = r["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            owner = getattr(apt, "owner_name", "") or ""
            for b in (r.get("carry_bills") or []):
                b_date = getattr(b, "due_date", None) or getattr(b, "created_at", None)
                desc = getattr(b, "description", None) or getattr(b, "desc", None) or ""
                status = getattr(b, "status", "") or ""
                remaining = float(getattr(b, "_remaining", None) if getattr(b, "_remaining", None) is not None else (getattr(b, "amount", 0) or 0))

                line = []
                if global_mode:
                    line.append(site_map.get(r["site_id"], ""))
                line += [apt_label, owner, _fmt_dt(b_date), desc, remaining, status, getattr(b, "id", None)]
                ws_cb.append(line)

        for col in range(1, len(cb_headers) + 1):
            ws_cb.column_dimensions[get_column_letter(col)].width = 22
        ws_cb.column_dimensions[get_column_letter(4 if global_mode else 3)].width = 14
        ws_cb.column_dimensions[get_column_letter(5 if global_mode else 4)].width = 48

        # -------------------------
        # 5) ✅ Devreden - Ödemeler
        # -------------------------
        ws_cp = wb.create_sheet("Devreden-Ödemeler")
        cp_headers = []
        if global_mode:
            cp_headers.append("Site")
        cp_headers += ["Daire", "Malik", "Tarih", "Yöntem", "Tutar", "Bill ID", "Payment ID"]
        ws_cp.append(cp_headers)

        for r in rows:
            apt = r["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            owner = getattr(apt, "owner_name", "") or ""
            for p in (r.get("carry_payments") or []):
                p_date = getattr(p, "payment_date", None)
                method = getattr(p, "method", None) or ""
                amount = float(getattr(p, "amount", 0) or 0)

                line = []
                if global_mode:
                    line.append(site_map.get(r["site_id"], ""))
                line += [apt_label, owner, _fmt_dt(p_date), method, amount, getattr(p, "bill_id", None), getattr(p, "id", None)]
                ws_cp.append(line)

        for col in range(1, len(cp_headers) + 1):
            ws_cp.column_dimensions[get_column_letter(col)].width = 20
        ws_cp.column_dimensions[get_column_letter(4 if global_mode else 3)].width = 14

        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)

        # dosya adı: "bu_daire" ise daha net olsun
        if apartment_id and rows:
            apt = rows[0]["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            filename = f"rapor_{apt_label}_{month_label or start.isoformat()}.xlsx"
        else:
            filename = f"rapor_{month_label or start.isoformat()}.xlsx"

        return send_file(
            bio,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @admin_bp.route("/ledger/export.pdf", methods=["GET"])
    @admin_required
    def apartment_ledger_export_pdf():
        """
        ✅ Türkçe karakter destekli (DejaVuSans) PDF
        ✅ "Bu Daire – PDF" (apartment_id) gelirse sadece o daire
        ✅ İşlemler bölümündeki TÜM detaylar (4 tablo) PDF'e düzenli/profesyonel şekilde basılır
        ✅ Çizgiler yazıyla çakışmaz + dikey kolon çizgileri + tutar kolonları sağ hizalı
        ✅ Sayfa numarası + rapor meta bilgisi
        """
        import os
        from datetime import datetime
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4

        admin_user = _get_current_admin()
        site_id, global_mode = _get_site_scope_for_admin(admin_user)

        if not global_mode and not site_id:
            flash("Bu işlem için bir siteye atanmış olmanız gerekiyor.", "error")
            return redirect(url_for("admin.dashboard"))

        month_label, start, end = _resolve_period_from_request()
        apartment_id = request.args.get("apartment_id", type=int)
        q = (request.args.get("q") or "").strip() or None

        rows, site_map = _ledger_build_dataset(site_id, global_mode, start, end, apartment_id, q)

        # -------------------------
        # Font (TR karakter)
        # -------------------------
        font_regular = "DejaVuSans"
        font_bold = "DejaVuSans-Bold"

        candidates = [
            ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ]

        # proje içi static/fonts (opsiyonel)
        try:
            base_dir = os.path.dirname(os.path.abspath(__file__))          # routes/
            proj_root = os.path.abspath(os.path.join(base_dir, ".."))      # proje kökü varsayımı
            candidates.append((
                os.path.join(proj_root, "static", "fonts", "DejaVuSans.ttf"),
                os.path.join(proj_root, "static", "fonts", "DejaVuSans-Bold.ttf"),
            ))
        except Exception:
            pass

        registered = False
        for reg_path, bold_path in candidates:
            if os.path.exists(reg_path) and os.path.exists(bold_path):
                try:
                    pdfmetrics.registerFont(TTFont(font_regular, reg_path))
                    pdfmetrics.registerFont(TTFont(font_bold, bold_path))
                    registered = True
                    break
                except Exception:
                    registered = False

        current_font = font_regular if registered else "Helvetica"
        bold_font = font_bold if registered else "Helvetica-Bold"
        current_size = 9

        # -------------------------
        # Helpers
        # -------------------------
        def _safe_text(s):
            return (s or "").replace("\n", " ").replace("\r", " ").strip()

        def _fmt_dt(d):
            if not d:
                return ""
            try:
                dd = getattr(d, "date", lambda: d)()
                return dd.strftime("%d.%m.%Y")
            except Exception:
                return str(d)

        def _money(x):
            # TR format: 1.234,56
            try:
                val = float(x or 0)
                s = f"{val:,.2f}"
                s = s.replace(",", "X").replace(".", ",").replace("X", ".")
                return f"₺{s}"
            except Exception:
                return f"₺{x or 0}"

        def _wrap_line(cnv, text, x, y, max_w, line_h):
            """
            Basit wrap: stringWidth ile max_w aşınca alt satıra geçer.
            Yazdırdığı son satırdan sonraki Y'yi döndürür.
            """
            txt = _safe_text(text)
            if not txt:
                cnv.drawString(x, y, "")
                return y - line_h

            words = txt.split()
            line = ""
            for w in words:
                cand = w if not line else (line + " " + w)
                if cnv.stringWidth(cand, current_font, current_size) <= max_w:
                    line = cand
                else:
                    cnv.drawString(x, y, line)
                    y -= line_h
                    line = w
            if line:
                cnv.drawString(x, y, line)
                y -= line_h
            return y

        # -------------------------
        # PDF Canvas
        # -------------------------
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        margin_x = 36
        y = height - 40
        period_end_inclusive = end - timedelta(days=1)

        # Sayfa no
        page_no = 1

        def _draw_footer():
            c.setFont(current_font, 8)
            c.drawRightString(width - margin_x, 18, f"Sayfa {page_no}")

        def _draw_page_header(is_continued=False):
            nonlocal y
            c.setFont(bold_font, 13 if not is_continued else 11)
            c.drawString(
                margin_x, height - 40,
                f"Daire Ekstresi Raporu ({start.strftime('%Y-%m-%d')} - {period_end_inclusive.strftime('%Y-%m-%d')})"
                if not is_continued else
                "Daire Ekstresi Raporu (Devam)"
            )

            y = height - 62
            c.setFont(current_font, 9)

            generated_at = datetime.now().strftime("%d.%m.%Y %H:%M")
            period_str = f"{start.strftime('%d.%m.%Y')} - {period_end_inclusive.strftime('%d.%m.%Y')}"
            filter_str = f"Daire ID: {apartment_id}" if apartment_id else "Tüm Daireler"
            if q:
                filter_str += f" | Arama: {q}"

            c.drawString(margin_x, y, f"Oluşturma: {generated_at}   |   Dönem: {period_str}   |   Filtre: {filter_str}")
            y -= 12
            c.drawString(margin_x, y, f"Toplam kayıt: {len(rows)}")
            y -= 10
            c.setLineWidth(0.8)
            c.line(margin_x, y, width - margin_x, y)
            y -= 14

        def _ensure_space(min_y=110):
            nonlocal y, page_no
            if y < min_y:
                _draw_footer()
                c.showPage()
                page_no += 1
                _draw_page_header(is_continued=True)

        def _section_title(txt):
            nonlocal y
            _ensure_space(140)
            c.setFont(bold_font, 10)
            c.drawString(margin_x, y, txt)
            y -= 10
            c.setFont(current_font, 9)

        def _draw_table(headers, rows_data, col_widths, right_align_cols=None):
            """
            headers: list[str]
            rows_data: list[list[str]]
            col_widths: list[int]
            right_align_cols: set[int] -> sağ hizalı kolon indexleri
            """
            nonlocal y
            right_align_cols = right_align_cols or set()

            _ensure_space(160)

            x0 = margin_x
            table_w = sum(col_widths)

            # --- helper: wrap satırlarına böl ---
            def _split_lines(text, max_w):
                txt = _safe_text(text)
                if not txt:
                    return [""]
                words = txt.split()
                lines = []
                cur = ""
                for w in words:
                    cand = w if not cur else cur + " " + w
                    if c.stringWidth(cand, current_font, current_size) <= max_w:
                        cur = cand
                    else:
                        if cur:
                            lines.append(cur)
                        cur = w
                if cur:
                    lines.append(cur)
                return lines

            pad_x = 4
            pad_top = 3
            line_h = 11
            min_row_h = 20

            # --- Header alanı ---
            header_top = y
            header_h = 18
            header_bottom = header_top - header_h

            c.setFont(bold_font, 9)
            x = x0
            for i, htxt in enumerate(headers):
                c.drawString(x + pad_x, header_top - 13, _safe_text(htxt))
                x += col_widths[i]

            # header çizgileri
            c.setLineWidth(0.8)
            c.line(x0, header_top, x0 + table_w, header_top)
            c.line(x0, header_bottom, x0 + table_w, header_bottom)

            # header dikey çizgiler
            c.setLineWidth(0.4)
            x = x0
            c.line(x, header_top, x, header_bottom)
            for wcol in col_widths:
                x += wcol
                c.line(x, header_top, x, header_bottom)

            y = header_bottom - 8
            c.setFont(current_font, 9)

            # --- Satırlar ---
            for rline in rows_data:
                _ensure_space(130)

                cell_lines = []
                max_lines = 1
                for i, cell in enumerate(rline):
                    lines = _split_lines(cell, col_widths[i] - (pad_x * 2))
                    cell_lines.append(lines)
                    max_lines = max(max_lines, len(lines))

                row_h = max(min_row_h, pad_top + (max_lines * line_h) + 6)
                row_top = y
                row_bottom = row_top - row_h

                # metinler
                x = x0
                for i, lines in enumerate(cell_lines):
                    ty = row_top - pad_top - 2
                    for ln in lines:
                        if i in right_align_cols:
                            c.drawRightString(x + col_widths[i] - pad_x, ty, ln)
                        else:
                            c.drawString(x + pad_x, ty, ln)
                        ty -= line_h
                    x += col_widths[i]

                # yatay alt çizgi
                c.setLineWidth(0.35)
                c.line(x0, row_bottom, x0 + table_w, row_bottom)

                # dikey çizgiler (satır)
                c.setLineWidth(0.25)
                x = x0
                c.line(x, row_top, x, row_bottom)
                for wcol in col_widths:
                    x += wcol
                    c.line(x, row_top, x, row_bottom)

                y = row_bottom - 6

            y -= 6

        # -------------------------
        # İlk sayfa header
        # -------------------------
        _draw_page_header(is_continued=False)

        # -------------------------
        # İçerik
        # -------------------------
        for r in rows:
            _ensure_space(170)

            apt = r["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            owner = getattr(apt, "owner_name", "") or ""
            site_name = site_map.get(r.get("site_id"), "") if global_mode else ""

            # Daire başlığı
            c.setFont(bold_font, 11)
            title_line = f"{(site_name + ' | ') if global_mode and site_name else ''}{apt_label} — {owner}"
            c.drawString(margin_x, y, _safe_text(title_line)[:120])
            y -= 12

            # Özet
            c.setFont(current_font, 9)
            summary = (
                f"Devreden Borç: {_money(r.get('carry_debt'))}   |   "
                f"Devreden Ödeme: {_money(r.get('carry_paid'))}   |   "
                f"Dönem Borç: {_money(r.get('billed'))}   |   "
                f"Dönem Ödeme: {_money(r.get('paid'))}   |   "
                f"Genel Açık: {_money(r.get('open_balance'))}"
            )
            y = _wrap_line(c, summary, margin_x, y, width - 2 * margin_x, 11)
            y -= 4

            # 1) Dönem Borçları
            _section_title("Dönem Borçları")
            bills = r.get("bills") or []
            if not bills:
                c.drawString(margin_x + 8, y, "Bu dönemde borç yok.")
                y -= 14
            else:
                table_rows = []
                for b in bills:
                    b_date = getattr(b, "due_date", None) or getattr(b, "created_at", None)
                    desc = getattr(b, "description", None) or getattr(b, "desc", None) or ""
                    status = getattr(b, "status", "") or ""
                    amount = float(getattr(b, "amount", 0) or 0)
                    table_rows.append([_fmt_dt(b_date), desc, _money(amount), status])

                _draw_table(
                    headers=["Tarih", "Açıklama", "Tutar", "Durum"],
                    rows_data=table_rows,
                    col_widths=[80, 260, 90, 90],
                    right_align_cols={2}
                )

            # 2) Dönem Ödemeleri
            _section_title("Dönem Ödemeleri")
            pays = r.get("payments") or []
            if not pays:
                c.drawString(margin_x + 8, y, "Bu dönemde ödeme yok.")
                y -= 14
            else:
                table_rows = []
                for p in pays:
                    p_date = getattr(p, "payment_date", None)
                    method = getattr(p, "method", None) or ""
                    amount = float(getattr(p, "amount", 0) or 0)
                    bill_id = getattr(p, "bill_id", None)
                    table_rows.append([_fmt_dt(p_date), method, _money(amount), f"{bill_id or ''}"])

                _draw_table(
                    headers=["Tarih", "Yöntem", "Tutar", "Bill"],
                    rows_data=table_rows,
                    col_widths=[80, 220, 90, 130],
                    right_align_cols={2}
                )

            # 3) Devreden Borçlar
            _section_title("Devreden Borçlar")
            carry_bills = r.get("carry_bills") or []
            if not carry_bills:
                c.drawString(margin_x + 8, y, "Devreden borç yok.")
                y -= 14
            else:
                table_rows = []
                for b in carry_bills:
                    b_date = getattr(b, "due_date", None) or getattr(b, "created_at", None)
                    desc = getattr(b, "description", None) or getattr(b, "desc", None) or ""
                    status = getattr(b, "status", "") or ""
                    remaining = float(getattr(b, "_remaining", None) if getattr(b, "_remaining", None) is not None else (getattr(b, "amount", 0) or 0))
                    table_rows.append([_fmt_dt(b_date), desc, _money(remaining), status])

                _draw_table(
                    headers=["Tarih", "Açıklama", "Kalan", "Durum"],
                    rows_data=table_rows,
                    col_widths=[80, 260, 90, 90],
                    right_align_cols={2}
                )

            # 4) Devreden Ödemeler
            _section_title("Devreden Ödemeler")
            carry_pays = r.get("carry_payments") or []
            if not carry_pays:
                c.drawString(margin_x + 8, y, "Devreden ödeme yok.")
                y -= 14
            else:
                table_rows = []
                for p in carry_pays:
                    p_date = getattr(p, "payment_date", None)
                    method = getattr(p, "method", None) or ""
                    amount = float(getattr(p, "amount", 0) or 0)
                    bill_id = getattr(p, "bill_id", None)
                    table_rows.append([_fmt_dt(p_date), method, _money(amount), f"{bill_id or ''}"])

                _draw_table(
                    headers=["Tarih", "Yöntem", "Tutar", "Bill"],
                    rows_data=table_rows,
                    col_widths=[80, 220, 90, 130],
                    right_align_cols={2}
                )

            # daire ayırıcı
            _ensure_space(90)
            c.setLineWidth(0.8)
            c.line(margin_x, y, width - margin_x, y)
            y -= 16

        # Son sayfa footer
        _draw_footer()

        c.save()
        buffer.seek(0)

        if apartment_id and rows:
            apt = rows[0]["apartment"]
            apt_label = f"{getattr(apt,'block','')}-{getattr(apt,'floor','')}-{getattr(apt,'number','')}".strip("-")
            filename = f"rapor_{apt_label}_{month_label or start.isoformat()}.pdf"
        else:
            filename = f"rapor_{month_label or start.isoformat()}.pdf"

        return send_file(buffer, as_attachment=True, download_name=filename, mimetype="application/pdf")
